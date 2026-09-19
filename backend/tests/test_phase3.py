import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)

def test_save_and_list_contracts():
    # Save a contract
    resp = client.post("/api/contracts", json={
        "title": "ভাড়া চুক্তি - ধানমন্ডি",
        "document_type": "tenancy_agreement",
        "language": "bn",
        "data": {
            "landlord_name": "আনিসুল হক",
            "rent_amount": "৩৫,০০০"
        }
    })
    assert resp.status_code == 200
    res_data = resp.json()
    assert res_data["status"] == "saved"
    cid = res_data["id"]
    print("[PASS] Contract Saved to SQLite DB. ID:", cid)

    # List contracts
    resp = client.get("/api/contracts")
    assert resp.status_code == 200
    items = resp.json()
    assert len(items) > 0
    print(f"[PASS] List Contracts Passed (Found {len(items)} items)")

    # Fetch single contract
    resp = client.get(f"/api/contracts/{cid}")
    assert resp.status_code == 200
    single = resp.json()
    assert single["data"]["landlord_name"] == "আনিসুল হক"
    print("[PASS] Single Contract Retrieval Passed")

    # Delete contract
    resp = client.delete(f"/api/contracts/{cid}")
    assert resp.status_code == 200
    assert resp.json()["status"] == "deleted"
    print("[PASS] Contract Deletion Passed")

def test_audit_uploaded_document():
    demo_contract = """
    This Agreement is entered into between ABC Corporation and XYZ Agency.
    Term: 1 year. Compensation: $5,000 per month.
    The Client may terminate this contract immediately at any time without notice.
    The Contractor is liable for all unlimited indemnities.
    """
    resp = client.post("/api/ai/audit-upload", json={
        "raw_text": demo_contract,
        "filename": "Client_Contract.txt"
    })
    assert resp.status_code == 200
    data = resp.json()
    assert "score" in data
    assert "risks_found" in data
    assert len(data["risks_found"]) > 0
    assert "missing_clauses" in data
    print(f"[PASS] Upload Audit Passed (Safety Score: {data['score']}/100, Risks: {len(data['risks_found'])})")

def test_signature_in_pdf():
    # Dummy 1x1 transparent PNG as data URL
    dummy_sig = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
    resp = client.post("/api/export/pdf", json={
        "document_type": "tenancy_agreement",
        "language": "bn",
        "data": {
            "landlord_name": "স্বাক্ষরকারী মালিক",
            "tenant_name": "স্বাক্ষরকারী ভাড়াটিয়া",
            "party1_signature": dummy_sig,
            "party2_signature": dummy_sig
        }
    })
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert len(resp.content) > 2000
    print(f"[PASS] E-Signature Embedded PDF Generated Successfully! (Bytes: {len(resp.content)})")

def test_extract_file_text():
    import io
    import docx

    # Test DOCX extraction
    d = docx.Document()
    d.add_paragraph("বাড়ি ভাড়ার চুক্তিপত্র টেস্ট")
    d.add_paragraph("ভাড়াটিয়া: জনাব রফিক")
    buf = io.BytesIO()
    d.save(buf)
    docx_bytes = buf.getvalue()

    resp = client.post(
        "/api/tools/extract-file-text",
        files={"file": ("test_doc.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert resp.status_code == 200
    res = resp.json()
    assert res["detected_format"] == "docx"
    assert "বাড়ি ভাড়ার চুক্তিপত্র টেস্ট" in res["extracted_text"]
    print(f"[PASS] DOCX Text Extracted Successfully! (Length: {res['character_count']})")

    # Test TXT extraction
    resp_txt = client.post(
        "/api/tools/extract-file-text",
        files={"file": ("test.txt", "অংশীদারি কারবার চুক্তি".encode("utf-8"), "text/plain")}
    )
    assert resp_txt.status_code == 200
    res_t = resp_txt.json()
    assert res_t["detected_format"] == "txt"
    assert "অংশীদারি কারবার চুক্তি" in res_t["extracted_text"]
    print(f"[PASS] TXT Text Extracted Successfully! (Length: {res_t['character_count']})")

if __name__ == "__main__":
    test_save_and_list_contracts()
    test_audit_uploaded_document()
    test_extract_file_text()
    test_signature_in_pdf()
    print("\n[SUCCESS] ALL PHASE 3 TESTS PASSED!")
