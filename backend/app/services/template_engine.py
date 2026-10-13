from pathlib import Path
from jinja2 import Environment, FileSystemLoader, select_autoescape
from typing import Dict, Any, List
from app.config import TEMPLATES_DIR
import io
from docx import Document

class TemplateEngine:
    def __init__(self):
        self.env = Environment(
            loader=FileSystemLoader(str(TEMPLATES_DIR)),
            autoescape=select_autoescape(['html', 'xml'])
        )
        self.templates_meta = {
            "tenancy_agreement": {
                "id": "tenancy_agreement",
                "title_bn": "বাড়ি / ফ্ল্যাট / দোকান ভাড়ার চুক্তি",
                "title_en": "Tenancy / Lease Agreement",
                "template_file": "tenancy_agreement.html",
                "category": "Property & Real Estate",
                "defaults_bn": {
                    "execution_date": "১২ সেপ্টেম্বর, ২০২৬",
                    "landlord_name": "মো: রফিকুল ইসলাম",
                    "landlord_father": "মরহুম আলহাজ্ব আব্দুল করিম",
                    "landlord_address": "বাড়ি নং-১২, রোড নং-০৫, সেক্টর-৩, উত্তরা, ঢাকা",
                    "landlord_nid": "১৯৮৫২৬৯২৪১৫১২৩৪৫৬",
                    "landlord_phone": "০১৭১১-২২৩৩৪৪",
                    "tenant_name": "মো: তানভীর আহমেদ",
                    "tenant_father": "মো: নূরুল হুদা",
                    "tenant_address": "গ্রাম: চরভদ্রাসন, থানা: কোতোয়ালী, জেলা: ফরিদপুর",
                    "tenant_nid": "১৯৯২২৬৯২৪১৫৯৮৭৬৫৪",
                    "tenant_phone": "০১৮১১-৫৫৬৬৭৭",
                    "property_address": "ফ্ল্যাট নং-৪বি (৪র্থ তলা), বাড়ি নং-১২, রোড নং-০৫, সেক্টর-৩, উত্তরা, ঢাকা-১২৩০",
                    "property_type": "আবাসিক ফ্ল্যাট (৩ বেড, ৩ বাথ, ড্রয়িং-ডাইনিং)",
                    "duration_months": "২৪",
                    "start_date": "০১ অক্টোবর, ২০২৬",
                    "rent_amount": "২৫,০০০",
                    "rent_in_words": "পঁচিশ হাজার টাকা মাত্র",
                    "payment_due_day": "৭",
                    "deposit_amount": "৭৫,০০০",
                    "notice_period_months": "২",
                    "utility_terms": "বিদ্যুৎ বিল, গ্যাস বিল, পানি বিল ও মাসিক কমন সার্ভিস চার্জ (৩,০০০ টাকা) ২য় পক্ষ প্রতি মাসে নির্ধারিত তারিখের মধ্যে পরিশোধ করিবেন।",
                    "witness1_name": "মো: কামরুল হাসান",
                    "witness1_address": "উত্তরা, ঢাকা",
                    "witness2_name": "সাজিদ মাহমুদ",
                    "witness2_address": "মিরপুর, ঢাকা",
                    "custom_clauses": [
                        {
                            "title": "নিরাপত্তা ও পারিবারিক অনুশাসন",
                            "text": "ভাড়াকৃত ফ্ল্যাটে কোনো প্রকার অসামাজিক, অবৈধ বা আইন পরিপন্থী কর্মকাণ্ড পরিচালনা করা যাইবে না এবং রাত ১১টার পর প্রধান গেটের নিরাপত্তা বজায় রাখিতে হইবে।"
                        }
                    ]
                }
            },
            "nda_agreement": {
                "id": "nda_agreement",
                "title_bn": "নন-ডিসক্লোজার ও বাণিজ্যিক গোপনীয়তা চুক্তি (NDA)",
                "title_en": "Non-Disclosure Agreement (NDA)",
                "template_file": "nda_agreement.html",
                "category": "Corporate & Tech",
                "defaults_bn": {
                    "execution_date": "১২ সেপ্টেম্বর, ২০২৬",
                    "disclosing_party_name": "টেকনোভেশন সফটওয়্যার লিমিটেড",
                    "disclosing_party_address": "লেভেল ৫, বনানী বাণিজ্যিক এলাকা, ঢাকা",
                    "disclosing_party_rep": "আসিফ ইকবাল (ব্যবস্থাপনা পরিচালক)",
                    "receiving_party_name": "ডাটাফ্লো অ্যানালিটিক্স ইনকর্পোরেশন",
                    "receiving_party_address": "রোড ১১, গুলশান ২, ঢাকা",
                    "receiving_party_rep": "ফাহিম জামান (চিফ টেকনোলজি অফিসার)",
                    "purpose": "প্রস্তাবিত ক্লাউড ইআরপি ও ফিনটেক প্ল্যাটফর্মের যৌথ প্রযুক্তিগত মূল্যায়ন ও কোড রিভিউ পরিচালনা করণ।",
                    "duration_years": "৩",
                    "custom_clauses": [
                        {
                            "title": "কর্মচারী নিয়োগ নিষেধাজ্ঞা (Non-Solicitation)",
                            "text": "চুক্তির মেয়াদকালীন সময়ে এবং চুক্তি অবসানের পরবর্তী ১ বছরের মধ্যে কোনো পক্ষ অন্য পক্ষের কোনো প্রকৌশলী বা কর্মীকে নিজ প্রতিষ্ঠানে প্রলুব্ধ বা নিয়োগ প্রদান করিতে পারিবে না।"
                        }
                    ]
                }
            },
            "freelance_contract": {
                "id": "freelance_contract",
                "title_bn": "ফ্রিল্যান্স সার্ভিস ও পরামর্শক চুক্তিপত্র",
                "title_en": "Freelance & Independent Contractor Agreement",
                "template_file": "freelance_contract.html",
                "category": "Freelancing & Services",
                "defaults_bn": {
                    "execution_date": "১২ সেপ্টেম্বর, ২০২৬",
                    "client_name": "নেক্সাস ডিজিটাল মার্কেটিং এজেন্সী",
                    "client_address": "ধানমন্ডি ২৭, ঢাকা",
                    "client_email": "client@nexusdigital.com",
                    "freelancer_name": "সাব্বির হোসেন",
                    "freelancer_title": "সিনিয়র ফুল-স্ট্যাক সফটওয়্যার ইঞ্জিনিয়ার",
                    "freelancer_address": "মিরপুর ডিওএইচএস, ঢাকা",
                    "freelancer_email": "sabbir.dev@gmail.com",
                    "scope_of_work": "Next.js, FastAPI এবং Stripe পেমেন্ট গেটওয়ে সম্বলিত আধুনিক SaaS ই-কমার্স ওয়েব অ্যাপ্লিকেশন তৈরি ও ক্লাউড ডেপ্লয়মেন্ট সম্পন্ন করা।",
                    "currency": "৳",
                    "total_amount": "১,৫০,০০০",
                    "payment_milestones": "প্রজেক্ট শুরুর পূর্বে ৩০% অগ্রিম (৳ ৪৫,০০০), আলফা রিলিজ ও UI অ্যাপ্রুভালে ৪০% (৳ ৬০,০০০), এবং সোর্স কোড ও ক্লাউড ডেপ্লয়মেন্টে বাকি ৩০% (৳ ৪৫,০০০) পরিশোধযোগ্য হইবে।",
                    "deadline": "১৫ নভেম্বর, ২০২৬",
                    "free_revisions": "৩",
                    "custom_clauses": [
                        {
                            "title": "ওয়ারেন্টি ও বাগ ফিক্সিং সহায়তা",
                            "text": "চূড়ান্ত ডেলিভারির পর পরবর্তী ৩০ দিন পর্যন্ত যেকোনো টেকনিক্যাল বাগ বা এরর ফ্রিল্যান্সার কোনো অতিরিক্ত পারিশ্রমিক ছাড়াই সংশোধন করিয়া দিবেন।"
                        }
                    ]
                }
            },
            "partnership_agreement": {
                "id": "partnership_agreement",
                "title_bn": "অংশীদারি কারবার চুক্তিপত্র (Partnership Deed)",
                "title_en": "Partnership Deed & Agreement",
                "template_file": "partnership_agreement.html",
                "category": "Business & Trade",
                "defaults_bn": {
                    "execution_date": "১২ সেপ্টেম্বর, ২০২৬",
                    "partner1_name": "আব্দুল্লাহ আল মামুন",
                    "partner1_father": "মরহুম সামসুল হক",
                    "partner1_address": "বনশ্রী, রামপুরা, ঢাকা",
                    "partner1_nid": "১৯৮৮২৬৯২৪১৫৫৬৭৮৯০",
                    "partner1_phone": "০১৭০০-১১২২৩৩",
                    "partner2_name": "মাহমুদুর রহমান",
                    "partner2_father": "মো: রেজাউল করিম",
                    "partner2_address": "খিলগাঁও, ঢাকা",
                    "partner2_nid": "১৯৯০২৬৯২৪১৫০৯৮৭৬৫",
                    "partner2_phone": "০১৮০০-৪৪৫৫৬৬",
                    "firm_name": "মেসার্স ব্লুমিং এগ্রো অ্যান্ড টেক সল্যুশনস",
                    "firm_address": "হাউজ ৪, ব্লক সি, বনশ্রী, ঢাকা",
                    "total_capital": "২০,০০,০০০",
                    "partner1_share": "৫০",
                    "partner2_share": "৫০",
                    "bank_operation": "উভয় অংশীদারের যৌথ স্বাক্ষরে (Joint Signatures)",
                    "notice_period_months": "৩",
                    "custom_clauses": [
                        {
                            "title": "নতুন অংশীদার গ্রহণ বিধি",
                            "text": "উভয় অংশীদারের সর্বসম্মত লিখিত সম্মতি ব্যতীত কোনো তৃতীয় ব্যক্তিকে এই ফার্মের অংশীদার হিসেবে অন্তর্ভুক্ত করা যাইবে না।"
                        }
                    ]
                }
            },
            "employment_agreement": {
                "id": "employment_agreement",
                "title_bn": "কর্মসংস্থান ও চাকরির চুক্তিপত্র (Employment Contract)",
                "title_en": "Employment Agreement & Appointment",
                "template_file": "employment_agreement.html",
                "category": "HR & Corporate",
                "defaults_bn": {
                    "execution_date": "১২ সেপ্টেম্বর, ২০২৬",
                    "company_name": "ইনোভেক্স ক্লাউড ল্যাবস লিমিটেড",
                    "company_address": "সফটওয়্যার টেকনোলজি পার্ক, কাওরান বাজার, ঢাকা",
                    "company_rep": "সাকিব আল হাসান (হেড অব এইচআর)",
                    "employee_name": "ফারহান কবির",
                    "employee_father": "মো: রফিকুল আলম",
                    "employee_address": "মোহাম্মদপুর, ঢাকা",
                    "employee_nid": "১৯৯৫২৬৯২৪১৫৭৮৯৪৫৬",
                    "employee_phone": "০১৯১১-২২৩৩৪৪",
                    "designation": "সফটওয়্যার কোয়ালিটি অ্যাসিউরেন্স ইঞ্জিনিয়ার",
                    "department": "প্রকৌশল ও পণ্য উন্নয়ন বিভাগ",
                    "salary_amount": "৬০,০০০",
                    "probation_months": "৩",
                    "notice_period_days": "৩০",
                    "custom_clauses": [
                        {
                            "title": "রিমোট ওয়ার্ক ও আইপি সুরক্ষা",
                            "text": "কর্মকর্তাকে প্রদত্ত প্রাতিষ্ঠানিক ল্যাপটপ ও ডিজিটাল ক্রেডেনশিয়াল শুধুমাত্র অফিশিয়াল কাজে ব্যবহৃত হইবে এবং কোনো সিকিউরিটি কোড তৃতীয় পক্ষের নিকট শেয়ার করা যাইবে না।"
                        }
                    ]
                }
            }
        }

    def list_templates(self) -> List[Dict[str, Any]]:
        result = []
        for t_id, meta in self.templates_meta.items():
            result.append({
                "id": t_id,
                "title_bn": meta["title_bn"],
                "title_en": meta["title_en"],
                "category": meta["category"],
                "defaults": meta["defaults_bn"]
            })
        return result

    def get_template_meta(self, doc_type: str) -> Dict[str, Any]:
        return self.templates_meta.get(doc_type, self.templates_meta["tenancy_agreement"])

    def render_html(self, doc_type: str, data: Dict[str, Any], language: str = "bn") -> str:
        meta = self.get_template_meta(doc_type)
        template = self.env.get_template(meta["template_file"])
        title = meta["title_bn"] if language == "bn" else meta["title_en"]
        if "ref_id" not in data:
            import hashlib
            core_str = f"{doc_type}_{data.get('execution_date', '')}_{data.get('landlord_name', '')}_{data.get('client_name', '')}_{data.get('partner1_name', '')}"
            short = hashlib.md5(core_str.encode('utf-8')).hexdigest()[:6].upper()
            data["ref_id"] = f"SLA-BD-{short}"

        if not data.get("qr_code_data"):
            from app.services.security_service import security_service
            verify_payload = f"/verify/{data['ref_id']}"
            data["qr_code_data"] = security_service.generate_qr_code_base64(verify_payload)

        return template.render(
            title=title,
            data=data,
            language=language
        )

    def generate_docx(self, doc_type: str, data: Dict[str, Any], language: str = "bn") -> io.BytesIO:
        import base64
        from docx.shared import Inches, Pt, RGBColor

        meta = self.get_template_meta(doc_type)
        doc = Document()
        title = meta["title_bn"] if language == "bn" else meta["title_en"]

        # Title
        heading = doc.add_heading(title, level=0)
        p_date = doc.add_paragraph()
        p_date.add_run(f"Execution Date / সম্পাদনের তারিখ: ").bold = True
        p_date.add_run(str(data.get('execution_date', '..................')))

        # Section 1: Parties
        doc.add_heading("1. Parties / পক্ষগণের বিবরণ", level=1)
        # Filter out internal/binary keys
        skip_keys = {"party1_signature", "party2_signature", "custom_clauses", "ref_id", "show_stamp"}
        for key, val in data.items():
            if key not in skip_keys and isinstance(val, (str, int, float)) and val:
                p = doc.add_paragraph()
                p.add_run(f"• {key.replace('_', ' ').title()}: ").bold = True
                p.add_run(str(val))

        # Section 2: Standard Legal Clauses
        doc.add_heading("2. Contract Terms & Clauses / মূল শর্তাবলি", level=1)
        if doc_type == "tenancy_agreement":
            clauses = [
                ("Term & Duration / চুক্তির মেয়াদ", f"মেয়াদ: {data.get('duration_months', '12')} মাস, কার্যকরের তারিখ: {data.get('start_date', '...')}" if language == "bn" else f"Duration: {data.get('duration_months', '12')} months, starting from {data.get('start_date', '...')}."),
                ("Monthly Rent / মাসিক ভাড়া", f"মাসিক ভাড়া: ৳ {data.get('rent_amount', '...')} টাকা ({data.get('rent_in_words', '')})। প্রতি মাসের {data.get('payment_due_day', '7')} তারিখের মধ্যে পরিশোধযোগ্য।" if language == "bn" else f"Monthly Rent: {data.get('rent_amount', '...')} BDT, payable by day {data.get('payment_due_day', '7')}."),
                ("Security Deposit / অগ্রিম জামানত", f"ফেরতযোগ্য অগ্রিম জামানত: ৳ {data.get('deposit_amount', '...')} টাকা।" if language == "bn" else f"Refundable security deposit: {data.get('deposit_amount', '...')} BDT."),
                ("Utility & Maintenance / ইউটিলিটি", str(data.get('utility_terms', 'বিদ্যুৎ, গ্যাস ও পানির বিল ভাড়াটিয়া নিজ দায়িত্বে পরিশোধ করিবেন।'))),
                ("Termination & Notice / চুক্তি অবসান", f"কমপক্ষে {data.get('notice_period_months', '2')} মাসের লিখিত নোটিশ সাপেক্ষে চুক্তি বাতিল করা যাইবে।" if language == "bn" else f"Termination requires at least {data.get('notice_period_months', '2')} month(s) written notice.")
            ]
        elif doc_type == "partnership_agreement":
            clauses = [
                ("Capital Contribution / মূলধন অনুপাত", f"মোট মূলধন: ৳ {data.get('total_capital', '...')} টাকা। ১ম অংশীদার: {data.get('partner1_share', '50')}%, ২য় অংশীদার: {data.get('partner2_share', '50')}%।" if language == "bn" else f"Total Capital: {data.get('total_capital', '...')} BDT. Partner 1: {data.get('partner1_share', '50')}%, Partner 2: {data.get('partner2_share', '50')}%."),
                ("Profit & Loss / লাভ-লোকসান বণ্টন", "অংশীদারদের মূলধনের অনুপাত অনুযায়ী লাভ-লোকসান বণ্টিত হইবে।" if language == "bn" else "Profits and losses shall be divided in proportion to capital contributions."),
                ("Bank Operations / ব্যাংক হিসাব পরিচালনা", f"ব্যাংক হিসাব: {data.get('bank_operation', 'উভয় অংশীদারের যৌথ স্বাক্ষরে পরিচালিত হইবে।')}" if language == "bn" else f"Bank Account: {data.get('bank_operation', 'Joint signatures')}."),
                ("Dissolution / কারবার অবসান", f"কারবার অবসানের জন্য কমপক্ষে {data.get('notice_period_months', '3')} মাসের লিখিত নোটিশ আবশ্যক।" if language == "bn" else f"Dissolution requires at least {data.get('notice_period_months', '3')} months notice.")
            ]
        elif doc_type == "freelance_contract":
            clauses = [
                ("Scope of Work / কাজের পরিধি", str(data.get('scope_of_work', 'চুক্তিতে উল্লেখিত নির্ধারিত সার্ভিস ও ডেলিভারেবলস সরবরাহকরণ।'))),
                ("Compensation / পারিশ্রমিক ও পেমেন্ট", f"মোট পারিশ্রমিক: {data.get('currency', '৳')} {data.get('total_amount', '...')}। {data.get('payment_milestones', '')}"),
                ("Deadline / ডেলিভারি সময়সীমা", f"চূড়ান্ত ডেলিভারির তারিখ: {data.get('deadline', 'নির্ধারিত তারিখ')}" if language == "bn" else f"Delivery deadline: {data.get('deadline', 'TBD')}."),
                ("Intellectual Property / স্বত্বাধিকার", "সম্পূর্ণ পারিশ্রমিক পরিশোধ সাপেক্ষে ডেলিভারেবলসের কপিরাইট গ্রাহকের অনুকূলে হস্তান্তরিত হইবে।" if language == "bn" else "All IP rights transfer to the client upon full payment.")
            ]
        elif doc_type == "employment_agreement":
            clauses = [
                ("Remuneration / বেতন ও সুবিধাদি", f"মাসিক বেতন: ৳ {data.get('salary_amount', '...')} টাকা।" if language == "bn" else f"Monthly salary: {data.get('salary_amount', '...')} BDT."),
                ("Probation / শিক্ষানবিসকাল", f"শিক্ষানবিসকাল: {data.get('probation_months', '3')} মাস।" if language == "bn" else f"Probation period: {data.get('probation_months', '3')} months."),
                ("Termination Notice / চাকরি অবসান", f"চাকরি অবসানের জন্য {data.get('notice_period_days', '30')} দিনের লিখিত নোটিশ প্রয়োজন।" if language == "bn" else f"Termination requires {data.get('notice_period_days', '30')} days notice.")
            ]
        else: # nda_agreement
            clauses = [
                ("Purpose of Disclosure / তথ্য বিনিময়ের উদ্দেশ্য", str(data.get('purpose', 'ব্যবসায়িক আলোচনা ও প্রযুক্তিগত মূল্যায়ন।'))),
                ("Confidentiality Term / গোপনীয়তার মেয়াদ", f"গোপনীয়তা রক্ষার মেয়াদ: {data.get('duration_years', '2')} বছর।" if language == "bn" else f"Duration: {data.get('duration_years', '2')} years."),
                ("Non-Disclosure Obligations / অ-প্রকাশ বাধ্যবাধকতা", "উভয় পক্ষ গোপনীয় তথ্যের সর্বোচ্চ সুরক্ষা বজায় রাখিতে অঙ্গীকারবদ্ধ।" if language == "bn" else "Strict standard of care required to maintain confidentiality.")
            ]

        for c_title, c_text in clauses:
            doc.add_heading(c_title, level=2)
            doc.add_paragraph(c_text)

        # Section 3: Special AI Clauses
        if "custom_clauses" in data and isinstance(data["custom_clauses"], list) and data["custom_clauses"]:
            doc.add_heading("3. Special AI Clauses / বিশেষ ধারাসমূহ", level=1)
            for c in data["custom_clauses"]:
                doc.add_heading(c.get("title", "Special Clause / বিশেষ ধারা"), level=2)
                doc.add_paragraph(c.get("text", ""))

        # Section 4: Signatures
        doc.add_heading("4. Signatures / স্বাক্ষর ও সম্মতি", level=1)
        p_sig_note = doc.add_paragraph("উভয় পক্ষ সুস্থ মস্তিষ্কে সজ্ঞানে এই চুক্তিতে স্বাক্ষর করিলেন।" if language == "bn" else "The parties hereto execute this agreement voluntarily.")

        def embed_signature(sig_key: str, party_name: str, label: str):
            p = doc.add_paragraph()
            p.add_run(f"\n{label} ({party_name or '...'}):\n").bold = True
            sig_val = data.get(sig_key)
            if sig_val and isinstance(sig_val, str) and sig_val.startswith("data:image"):
                try:
                    _, b64data = sig_val.split(",", 1)
                    img_bytes = base64.b64decode(b64data)
                    doc.add_picture(io.BytesIO(img_bytes), width=Inches(1.8))
                except Exception:
                    p.add_run("[Signed Digitally / ডিজিটাল স্বাক্ষর সংরক্ষিত]")
            else:
                p.add_run("স্বাক্ষর: ......................................................\n")

        p1_name = data.get('landlord_name') or data.get('client_name') or data.get('partner1_name') or data.get('company_name') or data.get('disclosing_party_name') or '১ম পক্ষ'
        p2_name = data.get('tenant_name') or data.get('freelancer_name') or data.get('partner2_name') or data.get('employee_name') or data.get('receiving_party_name') or '২য় পক্ষ'

        embed_signature('party1_signature', p1_name, "১ম পক্ষের স্বাক্ষর / First Party")
        embed_signature('party2_signature', p2_name, "২য় পক্ষের স্বাক্ষর / Second Party")

        # Footer
        p_foot = doc.add_paragraph()
        p_foot.add_run(f"\nSmartLegal AI • Verified Document | Ref: {data.get('ref_id', 'SLA-BD-CERTIFIED')}").italic = True

        target_stream = io.BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)
        return target_stream

template_engine = TemplateEngine()
