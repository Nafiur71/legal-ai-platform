import io
import logging
from typing import Tuple

logger = logging.getLogger("document_parser")

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None


def extract_text_from_bytes(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Extracts text from binary bytes of an uploaded file.
    Returns (clean_text, detected_format).
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext == "pdf":
        if not pypdf:
            raise RuntimeError("pypdf লাইব্রেরি ইনস্টল করা নেই।")
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for idx, page in enumerate(reader.pages):
            try:
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(text.strip())
            except Exception as e:
                logger.warning(f"Failed extracting text from page {idx}: {e}")
        extracted = "\n\n".join(pages_text)
        return extracted.strip(), "pdf"

    elif ext == "docx":
        if not docx:
            raise RuntimeError("python-docx লাইব্রেরি ইনস্টল করা নেই।")
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        extracted = "\n\n".join(paragraphs)
        return extracted.strip(), "docx"

    elif ext == "doc":
        try:
            decoded = file_bytes.decode("utf-8", errors="ignore").strip()
            clean = "".join(ch for ch in decoded if ch.isprintable() or ch in "\n\r\t ")
            return clean.strip(), "doc"
        except Exception:
            return "", "doc"

    else:
        for encoding in ["utf-8", "utf-8-sig", "cp1252", "latin-1"]:
            try:
                return file_bytes.decode(encoding).strip(), "txt"
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="ignore").strip(), "txt"
